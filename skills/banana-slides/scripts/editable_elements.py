from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt


def bounds(element, slide_width, slide_height):
    x, y, width, height = element["box"]
    return (
        round(x * slide_width), round(y * slide_height),
        round(width * slide_width), round(height * slide_height),
    )


def style_text(frame, text, element):
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = {
            "left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
        }[element.get("align", "left")]
        paragraph.space_before = paragraph.space_after = Pt(0)
        run = paragraph.add_run()
        run.text = line
        run.font.name = element.get("font", "Arial")
        run.font.size = Pt(element.get("font_size", 24))
        run.font.bold = element.get("bold", False)
        run.font.italic = element.get("italic", False)
        run.font.color.rgb = RGBColor.from_string(element.get("color", "222222"))


def add_text(slide, element, box):
    shape = slide.shapes.add_textbox(*box)
    style_text(shape.text_frame, element["text"], element)
    return shape


def add_shape(slide, element, box):
    kind = {"rect": MSO_SHAPE.RECTANGLE, "ellipse": MSO_SHAPE.OVAL,
            "arrow": MSO_SHAPE.RIGHT_ARROW}[element["shape"]]
    shape = slide.shapes.add_shape(kind, *box)
    if "fill" in element:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(element["fill"])
    else:
        shape.fill.background()
    if "stroke" in element:
        shape.line.color.rgb = RGBColor.from_string(element["stroke"])
        shape.line.width = Pt(element.get("stroke_width", 1))
    else:
        shape.line.fill.background()
    return shape


def add_table(slide, element, box):
    rows = element["rows"]
    shape = slide.shapes.add_table(len(rows), len(rows[0]), *box)
    for row_index, row in enumerate(rows):
        for column_index, text in enumerate(row):
            cell = shape.table.cell(row_index, column_index)
            style_text(cell.text_frame, text, element)
    return shape


def add_chart(slide, element, box):
    data = CategoryChartData()
    data.categories = element["categories"]
    for series in element["series"]:
        data.add_series(series["name"], series["values"])
    kind = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "line": XL_CHART_TYPE.LINE,
            "pie": XL_CHART_TYPE.PIE}[element["chart_type"]]
    shape = slide.shapes.add_chart(kind, *box, data)
    shape.chart.has_legend = element["chart_type"] == "pie" or len(element["series"]) > 1
    return shape


def add_native_element(slide, element, slide_width, slide_height):
    renderer = {"text": add_text, "shape": add_shape,
                "table": add_table, "chart": add_chart}[element["type"]]
    shape = renderer(slide, element, bounds(element, slide_width, slide_height))
    shape.name = element["id"]
    return shape
