import argparse
import io
import json
import sys
import warnings
from dataclasses import dataclass
from pathlib import Path

try:
    from jsonschema import Draft202012Validator
    from PIL import Image, ImageOps, UnidentifiedImageError
    from pptx import Presentation
    from pptx.util import Inches
except ImportError as error:
    raise SystemExit(
        "Missing dependency. Install scripts/requirements.txt with pip. "
        f"Details: {error}"
    ) from error


SCHEMA_PATH = Path(__file__).with_name("deck.schema.json")
SLIDE_WIDTH_INCHES = 10
SUPPORTED_FORMATS = {"PNG", "JPEG", "WEBP"}
MAX_MANIFEST_BYTES = 4 * 1024 * 1024
MAX_IMAGE_BYTES = 32 * 1024 * 1024


@dataclass(frozen=True)
class SlideImage:
    title: str
    notes: str
    png: bytes
    width: int
    height: int


def load_manifest(path: Path) -> dict:
    if path.stat().st_size > MAX_MANIFEST_BYTES:
        raise ValueError("Manifest exceeds 4 MiB.")
    data = json.loads(path.read_text(encoding="utf-8"))
    schema = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
    errors = tuple(Draft202012Validator(schema).iter_errors(data))
    if errors:
        details = "; ".join(
            f"{'.'.join(map(str, error.absolute_path)) or '$'}: {error.message}"
            for error in errors
        )
        raise ValueError(f"Invalid manifest: {details}")
    return data


def image_path(root: Path, value: str) -> Path:
    relative = Path(value)
    if relative.is_absolute() or ":" in value or "\\" in value:
        raise ValueError("Image must be a relative local path inside the deck folder.")
    resolved = (root / relative).resolve(strict=True)
    if not resolved.is_relative_to(root):
        raise ValueError("Image path escapes the deck folder.")
    if not resolved.is_file():
        raise ValueError("Image path is not a file.")
    if resolved.stat().st_size > MAX_IMAGE_BYTES:
        raise ValueError("Image exceeds 32 MiB.")
    return resolved


def read_slide(root: Path, data: dict, index: int) -> SlideImage:
    try:
        path = image_path(root, data["image"])
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(path) as original:
                if original.format not in SUPPORTED_FORMATS:
                    raise ValueError("Only PNG, JPEG, and WebP images are supported.")
                if getattr(original, "n_frames", 1) != 1:
                    raise ValueError("Animated images are not supported.")
                original.load()
                image = ImageOps.exif_transpose(original).convert("RGBA")
                buffer = io.BytesIO()
                image.save(buffer, format="PNG")
                return SlideImage(
                    data["title"], data.get("notes", ""), buffer.getvalue(),
                    image.width, image.height,
                )
    except (
        OSError, ValueError, UnidentifiedImageError,
        Image.DecompressionBombError, Image.DecompressionBombWarning,
    ) as error:
        raise ValueError(f"Slide {index + 1}: {error}") from error


def add_slide(prs, image: SlideImage) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    scale = min(prs.slide_width / image.width, prs.slide_height / image.height)
    width = round(image.width * scale)
    height = round(image.height * scale)
    picture = slide.shapes.add_picture(
        io.BytesIO(image.png),
        (prs.slide_width - width) // 2,
        (prs.slide_height - height) // 2,
        width=width,
        height=height,
    )
    picture.name = image.title
    slide.notes_slide.notes_text_frame.text = image.notes


def assemble(manifest: Path, output: Path) -> int:
    manifest = manifest.resolve(strict=True)
    output = output.absolute()
    if output.suffix.lower() != ".pptx":
        raise ValueError("Output must have a .pptx extension.")
    if output.exists() or output.is_symlink():
        raise ValueError("Output already exists. Choose a new output filename.")
    data = load_manifest(manifest)
    images = tuple(
        read_slide(manifest.parent, slide, index)
        for index, slide in enumerate(data["slides"])
    )
    ratio_width, ratio_height = map(int, data.get("aspect_ratio", "16:9").split(":"))
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_INCHES)
    prs.slide_height = Inches(SLIDE_WIDTH_INCHES * ratio_height / ratio_width)
    prs.core_properties.title = data["title"]
    prs.core_properties.author = "Banana Slides"
    for slide in images:
        add_slide(prs, slide)
    content = io.BytesIO()
    prs.save(content)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("xb") as target:
        target.write(content.getvalue())
    return len(images)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Assemble local slide images into an image-based PPTX."
    )
    parser.add_argument("manifest", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    try:
        count = assemble(args.manifest, args.output)
    except (OSError, ValueError, KeyError) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1
    print(json.dumps({"output": str(args.output), "slides": count}))
    return 0


if __name__ == "__main__":
    sys.exit(main())
