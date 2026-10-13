import argparse
import io
import json
import sys
from pathlib import Path

from assemble_pptx import load_manifest, read_slide, SLIDE_WIDTH_INCHES
from editable_elements import add_native_element, bounds
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches


SCHEMA_PATH = Path(__file__).with_name("editable-deck.schema.json")


def validate_element(element):
    x, y, width, height = element["box"]
    if x + width > 1.000001 or y + height > 1.000001:
        raise ValueError(f"Element {element['id']}: box exceeds slide bounds.")
    if element["type"] == "table":
        if len({len(row) for row in element["rows"]}) != 1:
            raise ValueError("Table rows must have the same number of cells.")
    if element["type"] == "text" and not element["text"].strip():
        raise ValueError("Text must not be blank.")
    if element["type"] != "chart":
        return
    labels = element["categories"] + [series["name"] for series in element["series"]]
    if any(label.startswith(("=", "{=")) for label in labels):
        raise ValueError("Chart labels cannot start with a spreadsheet formula.")
    for series in element["series"]:
        if len(series["values"]) != len(element["categories"]):
            raise ValueError("Chart values must match the category count.")
    if element["chart_type"] == "pie":
        values = element["series"][0]["values"]
        if len(element["series"]) != 1 or min(values) < 0 or sum(values) <= 0:
            raise ValueError("Pie charts require one nonnegative, nonzero series.")


def prepare_slide(root, slide, index):
    ids = [element["id"] for element in slide["elements"]]
    if len(ids) != len(set(ids)):
        raise ValueError(f"Slide {index + 1}: duplicate element ids.")
    if not any(e["type"] in {"text", "table", "chart"}
               for e in slide["elements"]):
        raise ValueError(f"Slide {index + 1}: editable content is required.")
    assets = {}
    for element in slide["elements"]:
        validate_element(element)
        if element["type"] == "image":
            assets[element["id"]] = read_slide(
                root, {"title": element["id"], "image": element["image"]}, index,
            )
    return assets


def add_image(slide, element, image, width, height):
    left, top, box_width, box_height = bounds(element, width, height)
    scale = min(box_width / image.width, box_height / image.height)
    image_width, image_height = round(image.width * scale), round(image.height * scale)
    picture = slide.shapes.add_picture(
        io.BytesIO(image.png), left + (box_width - image_width) // 2,
        top + (box_height - image_height) // 2,
        width=image_width, height=image_height,
    )
    picture.name = element["id"]


def assemble(manifest, output):
    manifest = manifest.resolve(strict=True)
    if output.suffix.lower() != ".pptx":
        raise ValueError("Output must have a .pptx extension.")
    if output.exists() or output.is_symlink():
        raise ValueError("Output already exists. Choose a new output filename.")
    data = load_manifest(manifest, SCHEMA_PATH)
    assets = tuple(prepare_slide(manifest.parent, slide, index)
                   for index, slide in enumerate(data["slides"]))
    prs = Presentation()
    ratio_width, ratio_height = map(int, data.get("aspect_ratio", "16:9").split(":"))
    prs.slide_width = Inches(SLIDE_WIDTH_INCHES)
    prs.slide_height = Inches(SLIDE_WIDTH_INCHES * ratio_height / ratio_width)
    prs.core_properties.title = data["title"]
    prs.core_properties.author = "Banana Slides"
    for data_slide, images in zip(data["slides"], assets):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(
            data_slide.get("background", "FFFFFF")
        )
        for element in data_slide["elements"]:
            if element["type"] == "image":
                add_image(slide, element, images[element["id"]],
                          prs.slide_width, prs.slide_height)
            else:
                add_native_element(slide, element, prs.slide_width, prs.slide_height)
        slide.notes_slide.notes_text_frame.text = data_slide.get("notes", "")
    content = io.BytesIO()
    prs.save(content)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("xb") as target:
        target.write(content.getvalue())
    return len(prs.slides)


def main():
    parser = argparse.ArgumentParser(description="Assemble editable slide elements.")
    parser.add_argument("manifest", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    try:
        count = assemble(args.manifest, args.output)
    except (OSError, ValueError, KeyError) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1
    print(json.dumps({"output": str(args.output), "slides": count, "editable": True}))
    return 0


if __name__ == "__main__":
    sys.exit(main())
