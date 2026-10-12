import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation


SCRIPT = (
    Path(__file__).resolve().parents[1]
    / "skills/banana-slides/scripts/assemble_pptx.py"
)


class AssemblePptxTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.deck = self.root / "deck"
        self.deck.mkdir()
        self.manifest = self.deck / "deck.json"
        self.output = self.root / "output" / "slides.pptx"
        Image.new("RGB", (160, 90), "red").save(self.deck / "red.png")
        Image.new("RGB", (90, 160), "blue").save(self.deck / "blue.png")
        self.data = {
            "title": "Test deck",
            "slides": [
                {"title": "Red", "image": "red.png", "notes": "中文 & <notes>"},
                {"title": "Blue", "image": "blue.png", "notes": "Second slide"},
            ],
        }

    def run_cli(self, data=None, output=None):
        self.manifest.write_text(
            json.dumps(self.data if data is None else data), encoding="utf-8"
        )
        return subprocess.run(
            [sys.executable, str(SCRIPT), str(self.manifest),
             "--output", str(output or self.output)],
            capture_output=True, text=True, check=False,
        )

    def assert_rejected(self, data, expected):
        result = self.run_cli(data)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn(expected, result.stderr)
        self.assertFalse(self.output.exists())

    def test_slide_order_aspect_ratio_and_notes(self):
        result = self.run_cli()
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout)["slides"], 2)
        deck = Presentation(self.output)
        self.assertEqual(len(deck.slides), 2)
        self.assertEqual(deck.slide_width / deck.slide_height, 16 / 9)
        self.assertEqual(deck.core_properties.title, "Test deck")
        for index, slide in enumerate(deck.slides):
            self.assertEqual(len(slide.shapes), 1)
            picture = slide.shapes[0]
            self.assertEqual(picture.name, self.data["slides"][index]["title"])
            self.assertEqual(picture.image.blob[:8], b"\x89PNG\r\n\x1a\n")
            self.assertEqual(
                slide.notes_slide.notes_text_frame.text,
                self.data["slides"][index]["notes"],
            )
        portrait = deck.slides[1].shapes[0]
        self.assertAlmostEqual(portrait.width / portrait.height, 90 / 160)
        self.assertGreater(portrait.left, 0)
        self.assertEqual(portrait.top, 0)

    def test_four_three_canvas(self):
        result = self.run_cli({**self.data, "aspect_ratio": "4:3"})
        self.assertEqual(result.returncode, 0, result.stderr)
        deck = Presentation(self.output)
        self.assertEqual(deck.slide_width / deck.slide_height, 4 / 3)

    def test_webp_input_becomes_embedded_png(self):
        Image.new("RGB", (160, 90), "green").save(self.deck / "green.webp")
        result = self.run_cli({
            **self.data, "slides": [{"title": "WebP", "image": "green.webp"}],
        })
        self.assertEqual(result.returncode, 0, result.stderr)
        picture = Presentation(self.output).slides[0].shapes[0]
        self.assertEqual(picture.image.content_type, "image/png")

    def test_exif_orientation_is_preserved(self):
        image = Image.new("RGB", (160, 90), "green")
        exif = Image.Exif()
        exif[274] = 6
        image.save(self.deck / "rotated.jpg", exif=exif)
        result = self.run_cli({
            **self.data, "slides": [{"title": "Exif", "image": "rotated.jpg"}],
        })
        self.assertEqual(result.returncode, 0, result.stderr)
        picture = Presentation(self.output).slides[0].shapes[0]
        self.assertEqual(picture.image.size, (90, 160))

    def test_missing_image_fails_without_partial_deck(self):
        (self.deck / "blue.png").unlink()
        self.assert_rejected(self.data, "Slide 2:")

    def test_corrupt_image_fails_without_partial_deck(self):
        (self.deck / "blue.png").write_text("not an image", encoding="utf-8")
        self.assert_rejected(self.data, "Slide 2:")

    def test_schema_rejects_invalid_inputs(self):
        cases = [
            ({**self.data, "slides": []}, "slides"),
            ({**self.data, "aspect_ratio": "0:0"}, "aspect_ratio"),
            ({**self.data, "unknown": True}, "unknown"),
            ({**self.data, "slides": [{"title": "Missing"}]}, "image"),
            ({**self.data, "title": False}, "title"),
        ]
        for data, expected in cases:
            with self.subTest(data=data):
                self.assert_rejected(data, expected)

    def test_absolute_paths_and_urls_are_rejected(self):
        for value in [str(self.deck / "red.png"), "https://example.com/a.png"]:
            with self.subTest(path=value):
                self.assert_rejected({
                    **self.data, "slides": [{"title": "Bad", "image": value}],
                }, "relative local path")

    def test_directory_escape_is_rejected(self):
        Image.new("RGB", (16, 9)).save(self.root / "outside.png")
        self.assert_rejected({
            **self.data, "slides": [{"title": "Bad", "image": "../outside.png"}],
        }, "escapes")

    def test_symlink_escape_is_rejected(self):
        Image.new("RGB", (16, 9)).save(self.root / "outside.png")
        (self.deck / "link.png").symlink_to(self.root / "outside.png")
        self.assert_rejected({
            **self.data, "slides": [{"title": "Bad", "image": "link.png"}],
        }, "escapes")

    def test_existing_output_is_not_overwritten(self):
        self.output.parent.mkdir()
        self.output.write_bytes(b"original")
        result = self.run_cli()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("already exists", result.stderr)
        self.assertEqual(self.output.read_bytes(), b"original")

    def test_invalid_json_has_actionable_error(self):
        self.manifest.write_text("{bad json", encoding="utf-8")
        result = subprocess.run(
            [sys.executable, str(SCRIPT), str(self.manifest),
             "--output", str(self.output)],
            capture_output=True, text=True, check=False,
        )
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("Error:", result.stderr)
        self.assertNotIn("Traceback", result.stderr)
        self.assertFalse(self.output.exists())


if __name__ == "__main__":
    unittest.main()
