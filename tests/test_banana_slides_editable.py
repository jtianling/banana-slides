import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData


SCRIPT = (Path(__file__).resolve().parents[1]
          / "skills/banana-slides/scripts/assemble_editable_pptx.py")


class EditableDeckTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.output = self.root / "editable.pptx"
        Image.new("RGB", (160, 90), "navy").save(self.root / "art.png")
        self.elements = [
            {"id": "art", "type": "image", "image": "art.png",
             "box": [0, 0, 1, 1]},
            {"id": "quote", "type": "text", "text": "中文可编辑\nEditable text",
             "box": [0.1, 0.1, 0.8, 0.2], "color": "FFFFFF", "bold": True},
            {"id": "divider", "type": "shape", "shape": "rect",
             "box": [0.1, 0.35, 0.8, 0.01], "fill": "FFFFFF"},
            {"id": "table", "type": "table", "rows": [["A", "B"], ["1", "2"]],
             "box": [0.1, 0.4, 0.35, 0.5], "font_size": 14},
            {"id": "chart", "type": "chart", "chart_type": "column",
             "categories": ["A", "B"],
             "series": [{"name": "Values", "values": [3, 7]}],
             "box": [0.5, 0.4, 0.4, 0.5]},
        ]

    def run_cli(self, elements=None, aspect="16:9"):
        manifest = self.root / "deck.json"
        manifest.write_text(json.dumps({
            "title": "Editable test", "aspect_ratio": aspect,
            "slides": [{"title": "First", "notes": "Source notes",
                        "elements": self.elements if elements is None else elements}],
        }), encoding="utf-8")
        return subprocess.run(
            [sys.executable, str(SCRIPT), str(manifest), "--output", str(self.output)],
            text=True, capture_output=True, check=False,
        )

    def test_native_objects_and_edit_save_reopen(self):
        result = self.run_cli()
        self.assertEqual(result.returncode, 0, result.stderr)
        prs = Presentation(self.output)
        slide = prs.slides[0]
        self.assertEqual(len(slide.shapes), 5)
        self.assertEqual(slide.shapes[1].text, "中文可编辑\nEditable text")
        self.assertTrue(slide.shapes[1].text_frame.paragraphs[0].runs[0].font.bold)
        self.assertTrue(slide.shapes[3].has_table)
        self.assertTrue(slide.shapes[4].has_chart)
        self.assertEqual(list(slide.shapes[4].chart.series[0].values), [3, 7])
        self.assertEqual(slide.notes_slide.notes_text_frame.text, "Source notes")
        slide.shapes[1].text = "修改后的文字"
        slide.shapes[3].table.cell(1, 1).text = "99"
        chart_data = CategoryChartData()
        chart_data.categories = ["A", "B"]
        chart_data.add_series("Changed", [9, 12])
        slide.shapes[4].chart.replace_data(chart_data)
        changed = self.root / "changed.pptx"
        prs.save(changed)
        reopened = Presentation(changed).slides[0]
        self.assertEqual(reopened.shapes[1].text, "修改后的文字")
        self.assertEqual(reopened.shapes[3].table.cell(1, 1).text, "99")
        self.assertEqual(list(reopened.shapes[4].chart.series[0].values), [9, 12])
        with ZipFile(self.output) as archive:
            self.assertTrue(any(n.startswith("ppt/embeddings/")
                                for n in archive.namelist()))

    def test_four_three_geometry(self):
        result = self.run_cli(aspect="4:3")
        self.assertEqual(result.returncode, 0, result.stderr)
        prs = Presentation(self.output)
        self.assertAlmostEqual(prs.slide_width / prs.slide_height, 4 / 3)
        self.assertEqual(prs.slides[0].shapes[1].left, round(prs.slide_width * 0.1))

    def test_invalid_elements_fail_without_output(self):
        cases = [
            [self.elements[0]],
            [{**self.elements[1], "text": " "}],
            [{**self.elements[1], "box": [0.9, 0, 0.2, 0.2]}],
            [{**self.elements[1], "color": "not-color"}],
            [{**self.elements[1], "font_size": float("nan")}],
            [self.elements[1], self.elements[1]],
            [{**self.elements[3], "rows": [["A", "B"], ["C"]]}],
            [{**self.elements[4], "categories": ["A"]}],
            [{**self.elements[4], "categories": ["=1+1", "B"]}],
            [{**self.elements[4],
              "series": [{"name": "{=1+1}", "values": [1, 2]}]}],
            [{**self.elements[4],
              "series": [{"name": "Huge", "values": [10 ** 400, 2]}]}],
            [{**self.elements[4], "chart_type": "pie",
              "series": [{"name": "Bad", "values": [-1, 0]}]}],
            [self.elements[1], {**self.elements[0], "image": "../outside.png"}],
        ]
        for elements in cases:
            with self.subTest(elements=elements):
                result = self.run_cli(elements)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("Error:", result.stderr)
                self.assertFalse(self.output.exists())

    def test_pie_categories_have_a_legend(self):
        result = self.run_cli([{**self.elements[4], "chart_type": "pie"}])
        self.assertEqual(result.returncode, 0, result.stderr)
        chart = Presentation(self.output).slides[0].shapes[0].chart
        self.assertTrue(chart.has_legend)

    def test_overflow_json_number_is_rejected(self):
        self.run_cli()
        self.output.unlink()
        manifest = self.root / "deck.json"
        data = manifest.read_text().replace('"values": [3, 7]', '"values": [1e309, 7]')
        self.assertIn("1e309", data)
        manifest.write_text(data)
        result = subprocess.run(
            [sys.executable, str(SCRIPT), str(manifest), "--output", str(self.output)],
            text=True, capture_output=True, check=False,
        )
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("Non-finite", result.stderr)
        self.assertNotIn("Traceback", result.stderr)
        self.assertFalse(self.output.exists())

    def test_does_not_overwrite_existing_file(self):
        self.output.write_bytes(b"original")
        result = self.run_cli()
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(self.output.read_bytes(), b"original")


if __name__ == "__main__":
    unittest.main()
